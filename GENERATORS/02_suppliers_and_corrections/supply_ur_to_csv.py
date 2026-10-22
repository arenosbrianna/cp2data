"""Reads makisushi_Jan_2026.xlsx, sheet SUPPLY-UR, into the supplier tables.
build_v4_supply_ur.py writes them to generated_v4/; running this file only prints a preview.

SUPPLY-UR is the only source for supplier names, products, quantities and prices. Every physical
row of the sheet is read from its cell values (as Excel last calculated them), whatever its
visibility: rows hidden by hand or by the AutoFilter, and grouped rows, are read like any other
(Excel stores a filtered-out row as a hidden row, so `row_visibility` covers both).

  suppliers                one row per vendor named in column B, ids in order of first appearance,
                           with the rows and products behind it. SUPPLY-UR has no delivery terms,
                           so terms_status is NOT_RECORDED.
  supply_unit_rates        one row per product row (every row from 6 down with a PRODUCT in C),
                           all columns A-S plus provenance, flags and the row's fate. Columns J and
                           N-S have no header; N-S are the owner's side calculations and keep their
                           column letter.
  term_assumptions         the SYNTHETIC delivery terms the simulation runs on (not from SUPPLY-UR).
  row_audit                every physical row 1..max_row with its visibility and fate.
  formula_audit            every formula cell: formula, cached value, meaning, where it went.

Normalization (source value kept next to it):
  * text is stripped of outer spaces: unit_source ' kg' -> unit 'kg' (rows 73, 74);
  * a product with no vendor in column B gets the vendor of the nearest row that has one (the row
    above on a tie); supplier_source records which row it came from. Rows 60-69 are HI-TOP by
    team instruction (SUPPLIER_OVERRIDE), not by that rule.
Row 77 'Kani - Sushi' is removed by team instruction (DROP_ROWS); row_audit records it as EXCLUDED.
A product name under more than one vendor is given to one vendor: the first row in sheet order
(is_chosen_supplier). The other rows stay in supply_unit_rates as REPRESENTED_AMBIGUOUS.
Left out (row_audit says why): the empty heading row 'BATTER' (79), the empty formula rows 89-90,
P4 (=H28*1.3) and N91:N93 (3.25, 1.35, =N91/2).
"""
import re
import warnings
from functools import lru_cache
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter, range_boundaries

ROOT = Path(__file__).resolve().parent
WORKBOOK, SHEET = ROOT / 'makisushi_Jan_2026.xlsx', 'SUPPLY-UR'
HEADER_ROW, FIRST_ROW = 4, 6
DROP_PRODUCTS = {'BATTER'}
# team decisions on rows the sheet leaves open (2026-10-05)
SUPPLIER_OVERRIDE = {r: 'HI-TOP' for r in range(60, 70)}   # blank-B rows between HI-TOP (59) and MY OWN (70)
DROP_ROWS = {77: "'Kani - Sushi' (HI-TOP) removed on team instruction 2026-10-05: Kani is supplied by "
                 "SHOKURYOHIN only (row 33); its N77 note 'ICHIBAN' is ignored with it"}
INSTRUCTION = 'team instruction 2026-10-05: rows 60-69 -> HI-TOP'
COLS = {  # sheet column -> csv column
    'A': 'col_a', 'B': 'supplier_name', 'C': 'product', 'D': 'group_code', 'E': 'quantity', 'F': 'unit',
    'G': 'detailed_quantity', 'H': 'unit_rate_php', 'I': 'amount_php', 'J': 'col_j',
    'K': 'wastage_fluctuation', 'L': 'net_qty', 'M': 'unit_rate_per_gms',
    'N': 'side_n', 'O': 'side_o', 'P': 'side_p', 'Q': 'side_q', 'R': 'side_r', 'S': 'side_s'}
HEADERS = {'B': 'SUPPLIER', 'C': 'PRODUCT', 'E': 'QUANTITY', 'F': 'UNIT', 'G': 'Detailed Quantity',
           'H': 'UNIT RATE', 'I': 'AMOUNT', 'K': 'Wastage/ Fluctuation', 'L': 'NET', 'M': 'Unit Rate per gms'}
NOT_RECORDED = 'NOT_RECORDED'

# SYNTHETIC delivery terms: SUPPLY-UR has no delivery data. Set from what each vendor sells
# and how Metro Manila suppliers of that kind usually deliver; not tuned to any KPI.
# They go to supplier_term_assumptions.csv, never to suppliers.csv.
#   lead_time_days          order placed -> delivered (>= 1)
#   lead_time_sd_days       day-to-day variation in that lead time
#   fill_rate               share of orders delivered in full
#   delivery_interval_days  1 = delivers daily, 2 = every other day, 7 = weekly
TERM_ASSUMPTIONS = {  # name: (lead_time_days, lead_time_sd_days, fill_rate, delivery_interval_days, basis)
    'BANRI':             (1, 0.3, 0.96, 2, 'fresh ramen / egg noodles (3-day life)'),
    'FIRST ENTERPRISES': (2, 0.5, 0.95, 7, 'bottled drinks by the case'),
    'FIRST REGAL':       (2, 0.5, 0.97, 7, 'cooking oil by the 18 L can'),
    'FRANCESCO':         (1, 0.3, 0.96, 1, 'salmon / tuna and other seafood, kept frozen (client Q&A)'),
    'HI-QRYS':           (1, 0.3, 0.98, 7, 'rice by the 25 kg sack'),
    'MJT GAS':           (1, 0.2, 0.99, 7, 'LPG tank refills'),
    'SHOKURYOHIN':       (2, 0.5, 0.95, 3, 'Japanese frozen, dry goods and sauces distributor'),
    'MY OWN':            (1, 0.2, 0.98, 2, 'beef, ground beef, ground pork'),
    'HI-TOP':            (1, 0.3, 0.96, 1, 'produce, tofu, mushroom (3-5 day life)'),
}
TERM_COLS = ['lead_time_days', 'lead_time_sd_days', 'fill_rate', 'delivery_interval_days']

# what each formula column computes (the formula text, not a business meaning)
FORMULA_MEANING = {
    'H': 'UNIT RATE built from other rows (Sushi Rice = 2 x Japanese Rice + Cane Vinegar + Sugar + 1)',
    'I': 'AMOUNT = UNIT RATE x QUANTITY',
    'L': 'NET = Detailed Quantity x (1 - Wastage/ Fluctuation)',
    'M': 'Unit Rate per gms = UNIT RATE / NET (0 if NET is 0 or blank)',
}


@lru_cache(maxsize=None)
def sheets():
    """(formulas, values) views of SUPPLY-UR."""
    warnings.filterwarnings('ignore')
    wf = openpyxl.load_workbook(WORKBOOK)[SHEET]
    wv = openpyxl.load_workbook(WORKBOOK, data_only=True)[SHEET]
    assert {c: wv[f'{c}{HEADER_ROW}'].value for c in HEADERS} == HEADERS, 'SUPPLY-UR header moved'
    return wf, wv


def clean(v):
    return v.strip() if isinstance(v, str) else v


def visibility():
    """Per-row visibility from the sheet's row attributes, and the AutoFilter range."""
    wf, _ = sheets()
    vis = {r: dict(row_visibility='HIDDEN' if wf.row_dimensions[r].hidden else 'VISIBLE',
                   outline_level=wf.row_dimensions[r].outlineLevel or 0)
           for r in range(1, wf.max_row + 1)}
    ref = wf.auto_filter.ref
    lo, hi = (range_boundaries(ref)[1], range_boundaries(ref)[3]) if ref else (0, -1)
    for r, v in vis.items():
        v['in_filter_range'] = lo <= r <= hi
    criteria = [fc.colId for fc in wf.auto_filter.filterColumn]
    return vis, ref, criteria


def is_formula(v):
    return isinstance(v, str) and v.startswith('=')


def nearest_supplier(rates):
    named = rates.dropna(subset=['supplier_name'])[['source_row', 'supplier_name']]
    names, sources, notes = [], [], []
    for r, name in zip(rates['source_row'], rates['supplier_name']):
        if isinstance(name, str):
            names.append(name); sources.append('SUPPLY-UR column B'); notes.append('')
            continue
        if r in SUPPLIER_OVERRIDE:
            names.append(SUPPLIER_OVERRIDE[r]); sources.append(INSTRUCTION); notes.append('')
            continue
        dist = (named['source_row'] - r).abs()
        best = named[dist == dist.min()].iloc[0]          # sorted by row, so a tie takes the row above
        above, below = named[named['source_row'] < r].tail(1), named[named['source_row'] > r].head(1)
        sides = [f"{s['source_row'].item()} {s['supplier_name'].item()}" for s in (above, below) if len(s)]
        agree = len(above) and len(below) and above['supplier_name'].item() == below['supplier_name'].item()
        names.append(best['supplier_name'])
        sources.append(f"nearest named row: {best['source_row']}")
        notes.append('' if agree else 'nearest named rows disagree or only one side: ' + ' / '.join(sides))
    return names, sources, notes


def product_rows():
    """Every product row of the sheet, raw and normalized, before supplier ids are given."""
    wf, wv = sheets()
    vis, _, _ = visibility()
    rows = []
    for r in range(FIRST_ROW, wv.max_row + 1):
        raw = {name: wv[f'{c}{r}'].value for c, name in COLS.items()}
        product = clean(raw['product'])
        if not (isinstance(product, str) and product) or product in DROP_PRODUCTS:
            continue
        rec = {name: clean(v) for name, v in raw.items()}
        changed = [f"{c} {raw[n]!r}->{rec[n]!r}" for c, n in COLS.items() if raw[n] != rec[n]]
        comments = [f"{c}{r}: {wf[f'{c}{r}'].comment.text.strip()}" for c in COLS if wf[f'{c}{r}'].comment]
        rows.append({'source_row': r, **vis[r], **rec,
                     'supplier_name_source': raw['supplier_name'], 'unit_source': raw['unit'],
                     'unit_rate_formula': wf[f'H{r}'].value if is_formula(wf[f'H{r}'].value) else None,
                     'net_qty_formula': wf[f'L{r}'].value if is_formula(wf[f'L{r}'].value) else None,
                     'normalization': '; '.join(changed), 'cell_comments': ' | '.join(comments)})
    rates = pd.DataFrame(rows)
    for c in ('col_a', 'col_j'):
        assert rates[c].isna().all(), f'{c} is not empty'
    return rates.drop(columns=['col_a', 'col_j'])


def data_flags(r):
    f = []
    if r.supplier_source == INSTRUCTION:
        f.append('SUPPLIER_SET_BY_TEAM_INSTRUCTION')
    elif r.supplier_source != 'SUPPLY-UR column B':
        f.append('SUPPLIER_FILLED_FROM_NEAREST_ROW')
    if r.supplier_fill_note:
        f.append('SUPPLIER_FILL_BOUNDARY')
    for col, flag in [('unit', 'MISSING_UNIT'), ('detailed_quantity', 'MISSING_DETAILED_QUANTITY'),
                      ('unit_rate_php', 'MISSING_UNIT_RATE'), ('amount_php', 'MISSING_AMOUNT'),
                      ('wastage_fluctuation', 'MISSING_WASTAGE_FLUCTUATION'),
                      ('unit_rate_per_gms', 'MISSING_UNIT_RATE_PER_GMS')]:
        if pd.isna(getattr(r, col)):
            f.append(flag)
    if r.unit_rate_per_gms == 0:
        f.append('UNIT_RATE_PER_GMS_ZERO')
    if pd.notna(r.wastage_fluctuation) and r.wastage_fluctuation < 0:
        f.append('NEGATIVE_WASTAGE_FLUCTUATION')
    if isinstance(r.net_qty_formula, str) and r.net_qty_formula.startswith('=1000*'):
        f.append('NET_USES_CONSTANT_1000_NOT_G')
    if isinstance(r.unit_rate_formula, str):
        f.append('UNIT_RATE_IS_FORMULA')
    if r.quantity != 1:
        f.append('QUANTITY_NOT_1')
    if r.normalization:
        f.append('TEXT_NORMALIZED')
    if r.duplicate_rows:
        f.append('PRODUCT_UNDER_SEVERAL_ROWS')
    if r.name_variant_of:
        f.append('PRODUCT_NAME_VARIANT')
    if isinstance(r.side_n, str):
        f.append('TEXT_IN_UNLABELLED_COLUMN_N')
    return ';'.join(f)


def build():
    """(suppliers, supply_unit_rates)."""
    rates = product_rows()
    names = list(dict.fromkeys(rates['supplier_name'].dropna()))
    assert set(names) == set(TERM_ASSUMPTIONS), f'term assumptions missing for {set(names) ^ set(TERM_ASSUMPTIONS)}'
    rates['supplier_name'], rates['supplier_source'], rates['supplier_fill_note'] = nearest_supplier(rates)
    rates = rates[~rates['source_row'].isin(DROP_ROWS)].reset_index(drop=True)   # after the fill: row 77 still names B
    sup =pd.DataFrame({'supplier_id': range(1, len(names) + 1), 'supplier_name': names})
    rates = rates.merge(sup, on='supplier_name', how='left')

    # one vendor per product name: the first row in sheet order
    first = rates.groupby('product')['source_row'].transform('min')
    rates['is_chosen_supplier'] = rates['source_row'] == first
    rows_of = rates.groupby('product')['source_row'].apply(list)
    rates['duplicate_rows'] = [';'.join(str(x) for x in rows_of[p] if x != r) for p, r in
                               zip(rates['product'], rates['source_row'])]
    prods = set(rates['product'])
    rates['name_variant_of'] = [';'.join(sorted(q for q in prods if q != p and (
        p.startswith(q + ' - ') or q.startswith(p + ' - ')))) for p in rates['product']]
    rates['wastage_fluctuation_semantics'] = 'UNKNOWN'
    rates['data_flags'] = [data_flags(r) for r in rates.itertuples()]

    def fate(r):
        if r.duplicate_rows:
            keep = rates.loc[(rates['product'] == r.product) & rates['is_chosen_supplier'], 'supplier_name'].item()
            return ('REPRESENTED_AMBIGUOUS', f"'{r.product}' is also on row(s) {r.duplicate_rows}; "
                    + ('chosen supplier (first row in sheet order)' if r.is_chosen_supplier
                       else f'not chosen: product given to {keep} (first row in sheet order)'))
        how = [x for x in (r.supplier_source if r.supplier_source != 'SUPPLY-UR column B' else '',
                           r.normalization) if x]
        return ('REPRESENTED_NORMALIZED', '; '.join(how)) if how else ('REPRESENTED_DIRECT', '')
    rates['row_fate'], rates['fate_detail'] = zip(*[fate(r) for r in rates.itertuples()])
    rates['source_sheet'] = SHEET

    for c in ('quantity', 'detailed_quantity'):
        assert (rates[c].dropna() % 1 == 0).all()
        rates[c] = rates[c].astype('Int64')
    lead = ['source_sheet', 'source_row', 'row_visibility', 'outline_level', 'in_filter_range',
            'supplier_id', 'supplier_name', 'supplier_name_source', 'supplier_source', 'supplier_fill_note',
            'product', 'group_code', 'quantity', 'unit_source', 'unit', 'detailed_quantity',
            'unit_rate_php', 'unit_rate_formula', 'amount_php', 'wastage_fluctuation',
            'wastage_fluctuation_semantics', 'net_qty', 'net_qty_formula', 'unit_rate_per_gms']
    tail = ['is_chosen_supplier', 'duplicate_rows', 'name_variant_of', 'normalization', 'cell_comments',
            'data_flags', 'row_fate', 'fate_detail']
    rates = rates[lead + [c for c in rates.columns if c not in lead + tail] + tail]
    return supplier_master(sup, rates), rates


def supplier_master(sup, rates):
    def rows(m):
        return ';'.join(str(x) for x in m['source_row'])
    out = []
    for s in sup.itertuples():
        m = rates[rates['supplier_id'] == s.supplier_id]
        named = m[m['supplier_source'] == 'SUPPLY-UR column B']
        vis = m['row_visibility'].value_counts()
        seen = {'VISIBLE', 'HIDDEN'} & set(vis.index)
        flags = sorted({f for fl in m['data_flags'] for f in fl.split(';') if f})
        incomplete = m[m['data_flags'].str.contains('MISSING_|UNIT_RATE_PER_GMS_ZERO')]
        conflict = m[m['data_flags'].str.contains('PRODUCT_UNDER_SEVERAL_ROWS|PRODUCT_NAME_VARIANT|SUPPLIER_FILL_BOUNDARY')]
        out.append(dict(
            supplier_id=s.supplier_id, supplier_name=s.supplier_name,
            supplier_name_source=';'.join(sorted(set(named['supplier_name_source']))),
            source_sheet=SHEET, supplier_source='SUPPLY-UR column B',
            source_rows_named=rows(named), source_rows_filled=rows(m.drop(named.index)),
            source_row_count=len(m), visible_row_count=int(vis.get('VISIBLE', 0)),
            hidden_row_count=int(vis.get('HIDDEN', 0)),
            appears_in='BOTH' if len(seen) == 2 else f'{seen.pop()}_ONLY',
            product_count=int(m['is_chosen_supplier'].sum()),
            products=';'.join(m.loc[m['is_chosen_supplier'], 'product']),
            products_not_chosen=';'.join(f'{p} (row {r})' for p, r in
                                         m.loc[~m['is_chosen_supplier'], ['product', 'source_row']].values),
            has_incomplete_rows=bool(len(incomplete)), incomplete_rows=rows(incomplete),
            has_conflicts=bool(len(conflict)), conflict_rows=rows(conflict),
            row_flags=';'.join(flags),
            terms_status=NOT_RECORDED, terms_source='none: SUPPLY-UR has no delivery terms'))
    return pd.DataFrame(out)


def term_assumptions(sup):
    t = pd.DataFrame([(n, *TERM_ASSUMPTIONS[n]) for n in sup['supplier_name']],
                     columns=['supplier_name', *TERM_COLS, 'basis'])
    t.insert(0, 'supplier_id', sup['supplier_id'].values)
    t['terms_status'] = 'SYNTHETIC_ASSUMPTION'
    t['terms_source'] = 'supply_ur_to_csv.TERM_ASSUMPTIONS (not SUPPLY-UR)'
    return t


def row_audit(rates):
    """Every physical row of SUPPLY-UR and what happened to it."""
    wf, wv = sheets()
    vis, _, _ = visibility()
    fate = rates.set_index('source_row')['row_fate']
    out = []
    for r in range(1, wf.max_row + 1):
        filled = [get_column_letter(c) for c in range(1, wf.max_column + 1) if wf.cell(r, c).value is not None]
        product = clean(wv[f'C{r}'].value)
        if r in fate.index:
            kind, f, note = 'PRODUCT', fate[r], 'supply_unit_rates row'
        elif r in DROP_ROWS:
            kind, f, note = 'PRODUCT', 'EXCLUDED', DROP_ROWS[r]
        elif not filled:
            kind, f, note = 'EMPTY', 'NO_DATA', ''
        elif r < HEADER_ROW:
            kind, f, note = 'TITLE', 'NOT_DATA', f'A{r} = {wv[f"A{r}"].value!r}'
        elif r == HEADER_ROW:
            kind, f, note = 'HEADER', 'NOT_DATA', 'column headers; P4 =H28*1.3 (318.5) is an unlabelled side calculation, left out'
        elif product in DROP_PRODUCTS:
            kind, f, note = 'SECTION_HEADING', 'EXCLUDED', f"'{product}' has no supplier, quantity or price and no product rows under it"
        elif set(filled) <= {'I', 'L', 'M'} and all(is_formula(wf[f'{c}{r}'].value) for c in filled):
            kind, f, note = 'FORMULA_ONLY', 'EXCLUDED', 'copied-down I/L/M formulas with no inputs (all 0); no supplier or product'
        elif set(filled) <= set('NOPQRS'):
            kind, f, note = 'SIDE_CALCULATION', 'EXCLUDED', 'unlabelled side value(s) ' + ', '.join(
                f"{c}{r}={wv[f'{c}{r}'].value!r}" for c in filled) + '; no supplier or product'
        else:
            raise ValueError(f'SUPPLY-UR row {r} ({filled}) has no rule')
        out.append(dict(excel_row=r, **vis[r], populated_columns=''.join(filled), row_kind=kind,
                        supplier_name=wv[f'B{r}'].value, product=product, fate=f, note=note))
    return pd.DataFrame(out)


def formula_audit():
    wf, wv = sheets()
    out = []
    for row in wf.iter_rows():
        for c in row:
            if not is_formula(c.value):
                continue
            col = c.column_letter
            if c.row in DROP_ROWS:
                where, meaning = 'not exported (row removed on team instruction)', FORMULA_MEANING.get(
                    col, 'owner side calculation in an unlabelled column')
            elif c.row >= FIRST_ROW and col in FORMULA_MEANING:
                where = (f'supply_unit_rates.{COLS[col]} (cached value)'
                         if clean(wv[f'C{c.row}'].value) not in (None, *DROP_PRODUCTS) else 'not exported (row has no product)')
                meaning = FORMULA_MEANING[col]
                if col == 'L' and c.value.startswith('=1000*'):
                    meaning = 'NET = 1000 x (1 - Wastage/ Fluctuation) (constant 1000, not Detailed Quantity)'
            else:
                where = (f'supply_unit_rates.{COLS[col]} (cached value)'
                         if c.row >= FIRST_ROW and clean(wv[f'C{c.row}'].value) else 'not exported')
                meaning = 'owner side calculation in an unlabelled column'
            out.append(dict(cell=c.coordinate, excel_row=c.row, column=col, header=HEADERS.get(col, ''),
                            formula=c.value, cached_value=wv[c.coordinate].value,
                            pattern=re.sub(r'(?<=[A-Z])' + str(c.row) + r'\b', '{row}', c.value),
                            meaning=meaning, exported_as=where))
    return pd.DataFrame(out)


def main():
    sup, rates = build()
    vis, ref, criteria = visibility()
    print(f"rows 1-{max(vis)}; hidden {sum(v['row_visibility'] == 'HIDDEN' for v in vis.values())}; "
          f"AutoFilter {ref}, criteria on columns {criteria or 'none'}")
    print(sup[['supplier_id', 'supplier_name', 'source_row_count', 'product_count', 'appears_in']].to_string(index=False))
    print(rates['row_fate'].value_counts().to_string())
    print(row_audit(rates)['fate'].value_counts().to_string())


if __name__ == '__main__':
    main()
